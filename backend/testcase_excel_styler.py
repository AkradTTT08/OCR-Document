import os
import re
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

def style_range(ws, cell_range, font=None, fill=None, border=None, alignment=None):
    if ":" in cell_range:
        cells = [c for row in ws[cell_range] for c in row]
    else:
        cells = [ws[cell_range]]
    for cell in cells:
        if font: cell.font = font
        if fill: cell.fill = fill
        if border: cell.border = border
        if alignment: cell.alignment = alignment

def safe_merge_and_style(ws, cell_range, value=None, font=None, fill=None, border=None, alignment=None):
    top_left = cell_range.split(":")[0]
    existing_ranges = [str(r) for r in ws.merged_cells.ranges]
    
    # If the exact range is already merged, don't re-merge
    if ":" in cell_range and cell_range not in existing_ranges:
        # Check if top-left is inside another conflicting merge
        for r in list(ws.merged_cells.ranges):
            if top_left in str(r) and str(r) != cell_range:
                try:
                    ws.unmerge_cells(str(r))
                except:
                    pass
        if value is not None:
            try:
                ws[top_left].value = value
            except:
                pass
        try:
            ws.merge_cells(cell_range)
        except Exception:
            pass
    elif ":" not in cell_range:
        if value is not None:
            try:
                ws[top_left].value = value
            except:
                pass
    else:
        if value is not None:
            try:
                ws[top_left].value = value
            except:
                pass

    style_range(ws, cell_range, font=font, fill=fill, border=border, alignment=alignment)

def calc_tc_row_height(step_text, obj_text="", exp_text="", data_text=""):
    lines_step = str(step_text or "").count("\n") + 1
    lines_step += len(str(step_text or "")) // 45
    lines_exp = str(exp_text or "").count("\n") + 1
    lines_exp += len(str(exp_text or "")) // 35
    lines_obj = str(obj_text or "").count("\n") + 1
    lines_obj += len(str(obj_text or "")) // 28
    lines_data = str(data_text or "").count("\n") + 1
    lines_data += len(str(data_text or "")) // 25
    max_lines = max(lines_step, lines_exp, lines_obj, lines_data, 1)
    return max(24, min(max_lines * 19, 180))

def beautify_workbook_file(file_path):
    if not os.path.exists(file_path):
        return False
    
    try:
        wb = openpyxl.load_workbook(file_path)
    except Exception as e:
        return False

    font_name = "Browallia New"

    # Color palette matching 69A template
    card_fill = PatternFill(start_color="CDEEFF", end_color="CDEEFF", fill_type="solid")
    header_fill = PatternFill(start_color="CDEEFF", end_color="CDEEFF", fill_type="solid")
    white_fill = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")
    dark_header_fill = PatternFill(start_color="7F7F7F", end_color="7F7F7F", fill_type="solid")

    thin_side = Side(style='thin', color='B0C4DE')
    dark_side = Side(style='thin', color='000000')
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    header_border = Border(left=dark_side, right=dark_side, top=dark_side, bottom=dark_side)

    title_banner_font = Font(name=font_name, size=26, bold=True, color="7F7F7F")
    title_font = Font(name=font_name, size=16, bold=True)
    section_white_font = Font(name=font_name, size=14, bold=True, color="FFFFFF")
    section_font = Font(name=font_name, size=14, bold=True, color="000000")
    header_font = Font(name=font_name, size=14, bold=True, color="000000")
    label_font = Font(name=font_name, size=14, bold=True, color="000000")
    value_font = Font(name=font_name, size=14, bold=False, color="000000")
    cell_font = Font(name=font_name, size=14, bold=False, color="000000")
    sheet_link_font = Font(name=font_name, size=14, bold=True, color="2E6433")

    align_center_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left_center = Alignment(horizontal="left", vertical="center", wrap_text=True)
    align_right_center = Alignment(horizontal="right", vertical="center")
    align_center_top = Alignment(horizontal="center", vertical="top", wrap_text=True)
    align_left_top = Alignment(horizontal="left", vertical="top", wrap_text=True)
    align_label = Alignment(horizontal="right", vertical="center")
    align_value = Alignment(horizontal="left", vertical="center", wrap_text=True)

    # 1. Introduction Sheet
    if "Introduction" in wb.sheetnames:
        ws = wb["Introduction"]
        try:
            ws.views.sheetView[0].showGridLines = True
        except:
            pass

        col_widths = {'A': 5.0, 'B': 10.0, 'C': 14.0, 'D': 14.0, 'E': 28.0, 'F': 24.0, 'G': 24.0, 'H': 14.0, 'I': 14.0}
        for col_l, w in col_widths.items():
            ws.column_dimensions[col_l].width = w

        for r in range(1, 5): ws.row_dimensions[r].height = 18
        intro_title = ws["A1"].value or "Introduction"
        safe_merge_and_style(ws, "A1:I4", value=intro_title, font=title_banner_font, fill=white_fill, alignment=align_center_center)

        sec_title = ws["A5"].value or "Introduction"
        ws.row_dimensions[5].height = 24
        safe_merge_and_style(ws, "A5:I5", value=sec_title, font=section_font, fill=header_fill, border=header_border, alignment=align_left_center)

        desc_text = ws["A6"].value or ""
        ws.row_dimensions[6].height = 22
        ws.row_dimensions[7].height = 22
        safe_merge_and_style(ws, "A6:I7", value=desc_text, font=value_font, fill=white_fill, border=cell_border, alignment=align_value)

        ws.row_dimensions[8].height = 12
        for r_num in (9, 10, 11):
            lbl_val = ws[f"A{r_num}"].value or ""
            c_val = ws[f"C{r_num}"].value or ws[f"B{r_num}"].value or ""
            ws.row_dimensions[r_num].height = 20
            safe_merge_and_style(ws, f"A{r_num}:B{r_num}", value=lbl_val, font=label_font, fill=white_fill, alignment=align_label)
            safe_merge_and_style(ws, f"C{r_num}:I{r_num}", value=c_val, font=value_font, fill=white_fill, alignment=align_value)

        struct_title = ws["A12"].value or "Structure of this workbook:"
        ws.row_dimensions[12].height = 24
        safe_merge_and_style(ws, "A12:I12", value=struct_title, font=section_font, fill=header_fill, border=header_border, alignment=align_left_center)

        sub_title = ws["A13"].value or "This workbook contains the following sheets and forms:"
        ws.row_dimensions[13].height = 20
        safe_merge_and_style(ws, "A13:I13", value=sub_title, font=value_font, fill=white_fill, alignment=align_left_center)

        ws.row_dimensions[14].height = 24
        safe_merge_and_style(ws, "B14:D14", value="Sheet Name", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)
        safe_merge_and_style(ws, "E14:I14", value="Detail", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)

        curr_r = 15
        while curr_r <= ws.max_row:
            s_name = ws[f"B{curr_r}"].value
            s_detail = ws[f"E{curr_r}"].value or ws[f"C{curr_r}"].value or ""
            if not s_name:
                break
            ws.row_dimensions[curr_r].height = 22
            safe_merge_and_style(ws, f"B{curr_r}:D{curr_r}", value=s_name, font=sheet_link_font, fill=white_fill, border=cell_border, alignment=align_left_center)
            safe_merge_and_style(ws, f"E{curr_r}:I{curr_r}", value=s_detail, font=cell_font, fill=white_fill, border=cell_border, alignment=align_left_center)
            curr_r += 1

        hist_start = curr_r + 2
        ws.row_dimensions[hist_start].height = 22
        safe_merge_and_style(ws, f"A{hist_start}:C{hist_start}", value="Template Change History:", font=section_white_font, fill=dark_header_fill, border=header_border, alignment=align_left_center)

        h_head_r = hist_start + 1
        ws.row_dimensions[h_head_r].height = 24
        safe_merge_and_style(ws, f"A{h_head_r}", value="Date", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)
        safe_merge_and_style(ws, f"B{h_head_r}", value="Version", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)
        safe_merge_and_style(ws, f"C{h_head_r}:D{h_head_r}", value="Author", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)
        safe_merge_and_style(ws, f"E{h_head_r}:I{h_head_r}", value="Detail", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)

        h_row_r = h_head_r + 1
        ws.row_dimensions[h_row_r].height = 24
        safe_merge_and_style(ws, f"A{h_row_r}", value="29/09/2026", font=cell_font, fill=white_fill, border=cell_border, alignment=align_center_center)
        safe_merge_and_style(ws, f"B{h_row_r}", value="v.2.0", font=cell_font, fill=white_fill, border=cell_border, alignment=align_center_center)
        safe_merge_and_style(ws, f"C{h_row_r}:D{h_row_r}", value="Lead QA Architect", font=cell_font, fill=white_fill, border=cell_border, alignment=align_center_center)
        safe_merge_and_style(ws, f"E{h_row_r}:I{h_row_r}", value="สร้าง System Test Case ตามข้อกำหนด SRS FoodSmile จัดระเบียบแยกตามโมดูลและเมนูระบบ", font=cell_font, fill=white_fill, border=cell_border, alignment=align_left_center)

    # 2. Document Change History Sheet
    if "Document Change History" in wb.sheetnames:
        ws = wb["Document Change History"]
        try: ws.views.sheetView[0].showGridLines = True
        except: pass
        col_w_hist = {'A': 2.0, 'B': 14.0, 'C': 12.0, 'D': 16.0, 'E': 16.0, 'F': 16.0, 'G': 30.0, 'H': 30.0, 'I': 30.0}
        for cl, w in col_w_hist.items(): ws.column_dimensions[cl].width = w

        for r in range(1, 5): ws.row_dimensions[r].height = 18
        safe_merge_and_style(ws, "B1:I4", value="Document Change History", font=title_banner_font, fill=white_fill, alignment=align_center_center)

        ws.row_dimensions[5].height = 22
        safe_merge_and_style(ws, "B5:E5", value="Document Change History:", font=section_white_font, fill=dark_header_fill, border=header_border, alignment=align_left_center)

        ws.row_dimensions[6].height = 24
        safe_merge_and_style(ws, "B6", value="Date", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)
        safe_merge_and_style(ws, "C6", value="Version", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)
        safe_merge_and_style(ws, "D6:F6", value="Prepared By", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)
        safe_merge_and_style(ws, "G6:I6", value="Detail", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)

        curr_r = 7
        while curr_r <= ws.max_row:
            d_val = ws[f"B{curr_r}"].value
            v_val = ws[f"C{curr_r}"].value
            p_val = ws[f"D{curr_r}"].value
            det_val = ws[f"G{curr_r}"].value or ws[f"E{curr_r}"].value or ""
            if not d_val and not v_val and not det_val:
                break
            safe_merge_and_style(ws, f"B{curr_r}", value=str(d_val or ""), font=cell_font, fill=white_fill, border=cell_border, alignment=align_center_top)
            safe_merge_and_style(ws, f"C{curr_r}", value=str(v_val or ""), font=cell_font, fill=white_fill, border=cell_border, alignment=align_center_top)
            safe_merge_and_style(ws, f"D{curr_r}:F{curr_r}", value=str(p_val or ""), font=cell_font, fill=white_fill, border=cell_border, alignment=align_center_top)
            safe_merge_and_style(ws, f"G{curr_r}:I{curr_r}", value=str(det_val or ""), font=cell_font, fill=white_fill, border=cell_border, alignment=align_left_top)

            lines = str(det_val).count("\n") + len(str(det_val)) // 70 + 1
            ws.row_dimensions[curr_r].height = max(24, lines * 18)
            curr_r += 1

    # 3. Glossary Sheet
    if "Glossary" in wb.sheetnames:
        ws = wb["Glossary"]
        try: ws.views.sheetView[0].showGridLines = True
        except: pass
        col_w_glo = {'A': 2.0, 'B': 35.0, 'C': 45.0, 'D': 45.0}
        for cl, w in col_w_glo.items(): ws.column_dimensions[cl].width = w

        for r in range(1, 5): ws.row_dimensions[r].height = 18
        safe_merge_and_style(ws, "B1:D4", value="Glossary", font=title_banner_font, fill=white_fill, alignment=align_center_center)

        ws.row_dimensions[5].height = 24
        safe_merge_and_style(ws, "B5", value="คำศัพท์ (Term)", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)
        safe_merge_and_style(ws, "C5:D5", value="คำอธิบาย (Definition)", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)

        curr_r = 6
        while curr_r <= ws.max_row:
            term_val = ws[f"B{curr_r}"].value
            def_val = ws[f"C{curr_r}"].value
            if not term_val and not def_val:
                break
            safe_merge_and_style(ws, f"B{curr_r}", value=str(term_val or ""), font=cell_font, fill=white_fill, border=cell_border, alignment=align_left_top)
            safe_merge_and_style(ws, f"C{curr_r}:D{curr_r}", value=str(def_val or ""), font=cell_font, fill=white_fill, border=cell_border, alignment=align_left_top)
            lines = str(def_val).count("\n") + len(str(def_val)) // 65 + 1
            ws.row_dimensions[curr_r].height = max(24, lines * 18)
            curr_r += 1

    # 4. Basic Test Sheet
    if "Basic Test" in wb.sheetnames:
        ws = wb["Basic Test"]
        try: ws.views.sheetView[0].showGridLines = True
        except: pass
        col_w_bas = {'A': 2.0, 'B': 10.0, 'C': 45.0, 'D': 30.0, 'E': 30.0, 'F': 25.0}
        for cl, w in col_w_bas.items(): ws.column_dimensions[cl].width = w

        for r in range(1, 5): ws.row_dimensions[r].height = 18
        safe_merge_and_style(ws, "B1:F4", value="Basic Test", font=title_banner_font, fill=white_fill, alignment=align_center_center)

        ws.row_dimensions[5].height = 22
        safe_merge_and_style(ws, "B5:C5", value="File List", font=section_white_font, fill=dark_header_fill, border=header_border, alignment=align_left_center)

        ws.row_dimensions[6].height = 24
        safe_merge_and_style(ws, "B6", value="No.", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)
        safe_merge_and_style(ws, "C6:D6", value="File Name", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)
        safe_merge_and_style(ws, "E6:F6", value="Location", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)

        ws.row_dimensions[7].height = 22
        safe_merge_and_style(ws, "B7", value="N/A", font=cell_font, fill=white_fill, border=cell_border, alignment=align_center_center)
        safe_merge_and_style(ws, "C7:D7", value="ไม่อยู่ในขอบเขตของเอกสารฉบับนี้ — ครอบคลุมเฉพาะ Functional/System Test Case", font=cell_font, fill=white_fill, border=cell_border, alignment=align_left_center)
        safe_merge_and_style(ws, "E7:F7", value="-", font=cell_font, fill=white_fill, border=cell_border, alignment=align_center_center)

        ws.row_dimensions[10].height = 22
        safe_merge_and_style(ws, "B10:C10", value="System Installation", font=section_white_font, fill=dark_header_fill, border=header_border, alignment=align_left_center)

        ws.row_dimensions[11].height = 24
        safe_merge_and_style(ws, "B11", value="No.", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)
        safe_merge_and_style(ws, "C11:F11", value="Setup Procedures", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)

        ws.row_dimensions[12].height = 22
        safe_merge_and_style(ws, "B12", value="N/A", font=cell_font, fill=white_fill, border=cell_border, alignment=align_center_center)
        safe_merge_and_style(ws, "C12:F12", value="ไม่อยู่ในขอบเขตของเอกสารฉบับนี้", font=cell_font, fill=white_fill, border=cell_border, alignment=align_left_center)

        ws.row_dimensions[15].height = 22
        safe_merge_and_style(ws, "B15:C15", value="System Configuration", font=section_white_font, fill=dark_header_fill, border=header_border, alignment=align_left_center)

        ws.row_dimensions[16].height = 24
        safe_merge_and_style(ws, "B16", value="No.", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)
        safe_merge_and_style(ws, "C16:F16", value="Configuration Procedures", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)

        ws.row_dimensions[17].height = 22
        safe_merge_and_style(ws, "B17", value="N/A", font=cell_font, fill=white_fill, border=cell_border, alignment=align_center_center)
        safe_merge_and_style(ws, "C17:F17", value="ไม่อยู่ในขอบเขตของเอกสารฉบับนี้", font=cell_font, fill=white_fill, border=cell_border, alignment=align_left_center)

    # 5. Execute Test Sheet
    if "Execute Test" in wb.sheetnames:
        ws = wb["Execute Test"]
        try: ws.views.sheetView[0].showGridLines = True
        except: pass
        col_w_exec = {'A': 2.0, 'B': 10.0, 'C': 30.0, 'D': 35.0, 'E': 35.0, 'F': 20.0, 'G': 20.0, 'H': 14.0, 'I': 16.0}
        for cl, w in col_w_exec.items(): ws.column_dimensions[cl].width = w

        for r in range(1, 5): ws.row_dimensions[r].height = 18
        safe_merge_and_style(ws, "B1:I4", value="Execute Test", font=title_banner_font, fill=white_fill, alignment=align_center_center)

        meta_exec = [
            (6, "Project Name :", ws["D6"].value or "FoodSmile"),
            (8, "Project ID:", ws["D8"].value or "69-PUT"),
            (10, "Project Release / Version :", ws["D10"].value or "2.0")
        ]
        for r_num, lbl_t, val_t in meta_exec:
            ws.row_dimensions[r_num].height = 22
            safe_merge_and_style(ws, f"C{r_num}", value=lbl_t, font=label_font, fill=card_fill, border=cell_border, alignment=align_right_center)
            safe_merge_and_style(ws, f"D{r_num}:F{r_num}", value=val_t, font=value_font, fill=white_fill, border=cell_border, alignment=align_value)

        ws.row_dimensions[13].height = 22
        safe_merge_and_style(ws, "B13:D13", value="System Test", font=section_white_font, fill=dark_header_fill, border=header_border, alignment=align_left_center)

        ws.row_dimensions[14].height = 32
        safe_merge_and_style(ws, "B14", value="REF #", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)
        safe_merge_and_style(ws, "C14", value="Module", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)
        safe_merge_and_style(ws, "D14", value="Function", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)
        safe_merge_and_style(ws, "E14", value="Pass\n(Meet criteria for evaluating)*", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)
        safe_merge_and_style(ws, "F14:G14", value="Remark", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)
        safe_merge_and_style(ws, "H14", value="จำนวน TC", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)
        safe_merge_and_style(ws, "I14", value="เวลาประมาณ (ชม.)", font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)

        curr_r = 15
        total_tc_sum = 0
        total_hrs_sum = 0.0
        while curr_r <= ws.max_row:
            ref_val = ws[f"B{curr_r}"].value
            mod_val = ws[f"C{curr_r}"].value
            if str(ref_val or '').strip() in ["รวม (Total)", "Total", "รวม"] or str(mod_val or '').strip() in ["รวม (Total)", "Total", "รวม"] or str(ws[f"D{curr_r}"].value or '').strip() in ["รวม (Total)", "Total", "รวม"]:
                break
            if not ref_val and not mod_val:
                break
            func_val = ws[f"D{curr_r}"].value or "-"
            pass_val = ws[f"E{curr_r}"].value or "ผ่านเกณฑ์การทดสอบตาม Requirement 100%"
            rem_val = ws[f"F{curr_r}"].value or "-"
            tc_c = int(ws[f"H{curr_r}"].value or 0)
            tc_hrs = float(ws[f"I{curr_r}"].value or 0.0)

            total_tc_sum += tc_c
            total_hrs_sum += tc_hrs

            ws.row_dimensions[curr_r].height = 24
            safe_merge_and_style(ws, f"B{curr_r}", value=str(ref_val), font=cell_font, fill=white_fill, border=cell_border, alignment=align_center_center)
            safe_merge_and_style(ws, f"C{curr_r}", value=str(mod_val), font=cell_font, fill=white_fill, border=cell_border, alignment=align_left_center)
            safe_merge_and_style(ws, f"D{curr_r}", value=str(func_val), font=cell_font, fill=white_fill, border=cell_border, alignment=align_left_center)
            safe_merge_and_style(ws, f"E{curr_r}", value=str(pass_val), font=cell_font, fill=white_fill, border=cell_border, alignment=align_center_center)
            safe_merge_and_style(ws, f"F{curr_r}:G{curr_r}", value=str(rem_val), font=cell_font, fill=white_fill, border=cell_border, alignment=align_center_center)
            safe_merge_and_style(ws, f"H{curr_r}", value=tc_c, font=cell_font, fill=white_fill, border=cell_border, alignment=align_center_center)
            safe_merge_and_style(ws, f"I{curr_r}", value=tc_hrs, font=cell_font, fill=white_fill, border=cell_border, alignment=align_center_center)
            curr_r += 1

        ws.row_dimensions[curr_r].height = 26
        safe_merge_and_style(ws, f"B{curr_r}:G{curr_r}", value="รวม (Total)", font=header_font, fill=header_fill, border=header_border, alignment=align_right_center)
        safe_merge_and_style(ws, f"H{curr_r}", value=total_tc_sum, font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)
        safe_merge_and_style(ws, f"I{curr_r}", value=round(total_hrs_sum, 2), font=header_font, fill=header_fill, border=header_border, alignment=align_center_center)

    # 6. Test Case Sheets
    for s_name in wb.sheetnames:
        if not s_name.lower().startswith("test case"):
            continue
        ws = wb[s_name]
        try: ws.views.sheetView[0].showGridLines = True
        except: pass

        # Check column structure: is it 11 columns (col 8 == Remark) or 13 columns?
        is_11_cols = (str(ws.cell(17, 8).value or '').strip() == "Remark")

        # Extract all existing test cases into memory before rebuilding
        raw_tcs = []
        curr_r = 18
        while curr_r <= ws.max_row:
            tc_id = ws[f"B{curr_r}"].value
            tc_obj = ws[f"C{curr_r}"].value
            if not tc_id and not tc_obj:
                break
            step_v = ws[f"D{curr_r}"].value or ""
            data_v = ws[f"E{curr_r}"].value or "-"
            type_v = ws[f"F{curr_r}"].value or "Positive"
            exp_v = ws[f"G{curr_r}"].value or ""
            
            if is_11_cols:
                act_v = ""
                status_v = ""
                rem_v = ws[f"H{curr_r}"].value or ""
                auto_v = ws[f"I{curr_r}"].value or "TRUE"
                req_v_row = ws[f"J{curr_r}"].value or "-"
                plat_v = ws[f"K{curr_r}"].value or "Web Application"
                upd_v = ws[f"L{curr_r}"].value or tester_v
            else:
                act_v = ws[f"H{curr_r}"].value or ""
                status_v = ws[f"I{curr_r}"].value or ""
                rem_v = ws[f"J{curr_r}"].value or ""
                auto_v = ws[f"K{curr_r}"].value or "TRUE"
                req_v_row = ws[f"L{curr_r}"].value or "-"
                plat_v = ws[f"M{curr_r}"].value or "Web Application"
                upd_v = ws[f"N{curr_r}"].value or tester_v

            raw_tcs.append({
                "id": tc_id, "obj": tc_obj, "step": step_v, "data": data_v,
                "type": type_v, "exp": exp_v, "act": act_v, "status": status_v,
                "rem": rem_v, "auto": auto_v, "req": req_v_row, "plat": plat_v, "upd": upd_v
            })
            curr_r += 1

        col_widths = {
            'A': 2.0,
            'B': 22.0, # Test Case ID
            'C': 26.0, # Test Case Objective
            'D': 45.0, # Test Step
            'E': 26.0, # Test Data
            'F': 14.0, # Test Type
            'G': 32.0, # Expected Result
            'H': 32.0, # Actual Result
            'I': 16.0, # Status (Pass/Fail)
            'J': 22.0, # Remark
            'K': 12.0, # Automate
            'L': 16.0, # Req No.
            'M': 16.0, # Platforms
            'N': 18.0  # Updated By
        }
        for cl, w in col_widths.items(): ws.column_dimensions[cl].width = w

        for r in range(1, 5): ws.row_dimensions[r].height = 18
        safe_merge_and_style(ws, "B1:N4", value="Test Case", font=title_banner_font, fill=white_fill, alignment=align_center_center)

        # Helper to extract clean meta value without label prefix
        def extract_clean_val(cands, default_v=""):
            for v in cands:
                s = str(v or '').strip()
                if s and not s.endswith(":") and "Date" not in s and "Function" not in s and "Module" not in s:
                    return s
            return default_v

        proj_name_v = extract_clean_val([ws["D6"].value, ws["E6"].value], "FoodSmile")
        proj_id_v = extract_clean_val([ws["D8"].value, ws["E8"].value], "69-PUT")
        tester_v = extract_clean_val([ws["D10"].value, ws["E10"].value], "Lead QA Architect")
        ver_v = extract_clean_val([ws["D12"].value, ws["E12"].value], "2.0")

        create_d_v = extract_clean_val([ws["I6"].value, ws["H6"].value, ws["G6"].value], "29/09/2026")
        start_d_v = extract_clean_val([ws["I8"].value, ws["H8"].value, ws["G8"].value], "29/09/2026")
        finish_d_v = extract_clean_val([ws["I10"].value, ws["H10"].value, ws["G10"].value], "29/09/2026")
        mod_func_v = extract_clean_val([ws["I12"].value, ws["H12"].value, ws["G12"].value], s_name)
        req_v = ws["E15"].value or ws["G15"].value or "REQ0001-REQ0050"

        # 1. Safely remove existing merged ranges in rows 5..13
        from openpyxl.cell.cell import Cell
        valid_ranges = []
        for m in ws.merged_cells.ranges:
            rows = [int(x) for x in re.findall(r'\d+', str(m))]
            if not any(r in range(5, 14) for r in rows):
                valid_ranges.append(m)
        ws.merged_cells.ranges = valid_ranges

        # 2. Reset all cells in B5:N13 to normal Cell instances to avoid MergedCell read-only issues
        for r in range(5, 14):
            for c in range(2, 15):
                ws._cells[(r, c)] = Cell(ws, row=r, column=c)

        # 3. Clear B5:N13 and set uniform corporate card fill and outer borders
        for r_idx in range(5, 14):
            ws.row_dimensions[r_idx].height = 20
            for c_idx in range(2, 15):
                cell = ws.cell(r_idx, c_idx)
                cell.value = None
                cell.fill = card_fill
                t_s = dark_side if r_idx == 5 else None
                b_s = dark_side if r_idx == 13 else None
                l_s = dark_side if c_idx == 2 else None
                r_s = dark_side if c_idx == 14 else None
                cell.border = Border(top=t_s, bottom=b_s, left=l_s, right=r_s)

        # 4. Perfectly symmetrical 4-row layout across B to N:
        # Left: label in C (align right), value box in D:F (merged white box)
        # Right: label in G:H (merged, align right), value box in I:M (merged white box)
        card_rows = [
            (6, "Project Name :", proj_name_v, "Create Date :", create_d_v),
            (8, "Project ID:", proj_id_v, "Start Test Date :", start_d_v),
            (10, "Tester Name :", tester_v, "Finish Test Date :", finish_d_v),
            (12, "Project Release / Version :", ver_v, "Module / Function: ", mod_func_v)
        ]
        for r_num, l_lbl, l_val, r_lbl, r_val in card_rows:
            # Left label in C
            ws.cell(r_num, 3, l_lbl).font = label_font
            ws.cell(r_num, 3).alignment = align_label

            # Left white box in D:F
            ws.merge_cells(f"D{r_num}:F{r_num}")
            ws[f"D{r_num}"] = l_val
            for c in ws[f"D{r_num}:F{r_num}"][0]:
                c.font = value_font
                c.fill = white_fill
                c.border = cell_border
                c.alignment = align_value

            # Right label in G:H
            ws.merge_cells(f"G{r_num}:H{r_num}")
            ws[f"G{r_num}"] = r_lbl
            for c in ws[f"G{r_num}:H{r_num}"][0]:
                c.font = label_font
                c.alignment = align_label

            # Right white box in I:M
            ws.merge_cells(f"I{r_num}:M{r_num}")
            ws[f"I{r_num}"] = r_val
            for c in ws[f"I{r_num}:M{r_num}"][0]:
                c.font = value_font
                c.fill = white_fill
                c.border = cell_border
                c.alignment = align_value

        ws.row_dimensions[14].height = 14
        ws.row_dimensions[15].height = 24
        safe_merge_and_style(ws, "B15:D15", value=" FUNCTIONAL REQUIREMENTS (Requirements No.) :", font=label_font, alignment=align_right_center)
        safe_merge_and_style(ws, "E15:H15", value=req_v, font=label_font, alignment=align_value)

        ws.row_dimensions[16].height = 12
        ws.row_dimensions[17].height = 32

        tc_headers = [
            ("Test Case ID", 2),
            ("Test Case Objective", 3),
            ("Test Step", 4),
            ("Test Data", 5),
            ("Test Type", 6),
            ("Expected Result", 7),
            ("Actual Result", 8),
            ("Status", 9),
            ("Remark", 10),
            ("Automate", 11),
            ("Req No.", 12),
            ("Platforms", 13),
            ("Updated By", 14)
        ]
        for h_name, col_i in tc_headers:
            cell = ws.cell(17, col_i, h_name)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = align_center_center
            cell.border = header_border

        tc_row_idx = 18
        for tc_item in raw_tcs:
            step_v = tc_item["step"]
            tc_obj = tc_item["obj"]
            exp_v = tc_item["exp"]
            data_v = tc_item["data"]

            ws.row_dimensions[tc_row_idx].height = calc_tc_row_height(step_v, tc_obj, exp_v, data_v)

            row_vals = [
                (2, tc_item["id"], align_center_top),
                (3, tc_obj, align_left_top),
                (4, step_v, align_left_top),
                (5, data_v, align_left_top),
                (6, tc_item["type"], align_center_top),
                (7, exp_v, align_left_top),
                (8, tc_item["act"], align_left_top),
                (9, tc_item["status"], align_center_top),
                (10, tc_item["rem"], align_left_top),
                (11, str(tc_item["auto"]).upper(), align_center_top),
                (12, tc_item["req"], align_center_top),
                (13, tc_item["plat"], align_center_top),
                (14, tc_item["upd"], align_center_top)
            ]
            for col_i, val, align_style in row_vals:
                c_node = ws.cell(tc_row_idx, col_i, val)
                c_node.font = cell_font
                c_node.alignment = align_style
                c_node.border = cell_border

            tc_row_idx += 1

    wb.save(file_path)
    return True
